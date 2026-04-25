"""Export functions for table extraction annotations.

Converts the annotation dict produced by ``pipeline.run_pipeline`` into
HTML table snippets, CSV strings, JSON, or Excel workbooks.
"""
from __future__ import annotations

import csv
import io
import json
from html import escape


def export_html(annotation: dict) -> str:
    """Convert annotation to self-contained HTML ``<table>`` snippet(s).

    One ``<table>`` per detected table.  Spanning cells use ``rowspan`` /
    ``colspan`` attributes on ``<td>`` elements.
    """
    html_parts: list[str] = []

    for table in annotation.get("tables", []):
        cells = table.get("cells", [])
        if not cells:
            continue

        max_row = max(c["row"] + c["row_span"] for c in cells)
        max_col = max(c["col"] + c["col_span"] for c in cells)

        # Positions covered by a span's non-anchor cells
        covered: set[tuple[int, int]] = set()
        for c in cells:
            if c["row_span"] > 1 or c["col_span"] > 1:
                for ri in range(c["row"], c["row"] + c["row_span"]):
                    for ci in range(c["col"], c["col"] + c["col_span"]):
                        if (ri, ci) != (c["row"], c["col"]):
                            covered.add((ri, ci))

        cell_map = {(c["row"], c["col"]): c for c in cells}

        html_parts.append('<table border="1">')
        for r in range(max_row):
            html_parts.append("  <tr>")
            for col in range(max_col):
                if (r, col) in covered:
                    continue
                cell = cell_map.get((r, col))
                if cell is None:
                    html_parts.append("    <td></td>")
                    continue
                attrs: list[str] = []
                if cell["row_span"] > 1:
                    attrs.append(f'rowspan="{cell["row_span"]}"')
                if cell["col_span"] > 1:
                    attrs.append(f'colspan="{cell["col_span"]}"')
                attr_str = (" " + " ".join(attrs)) if attrs else ""
                text = escape(cell.get("text", ""))
                html_parts.append(f"    <td{attr_str}>{text}</td>")
            html_parts.append("  </tr>")
        html_parts.append("</table>")

    return "\n".join(html_parts)


def export_csv(annotation: dict, table_id: int) -> str:
    """Convert a single table to a UTF-8 CSV string.

    Spanning cells have their text in the top-left (anchor) cell; other
    covered positions are left empty.

    Raises ``ValueError`` if *table_id* is not found.
    """
    table = None
    for t in annotation.get("tables", []):
        if t["table_id"] == table_id:
            table = t
            break
    if table is None:
        raise ValueError(f"Table ID {table_id} not found in annotation.")

    cells = table.get("cells", [])
    if not cells:
        return ""

    max_row = max(c["row"] + c["row_span"] for c in cells)
    max_col = max(c["col"] + c["col_span"] for c in cells)

    grid = [["" for _ in range(max_col)] for _ in range(max_row)]
    for c in cells:
        grid[c["row"]][c["col"]] = c.get("text", "")

    buf = io.StringIO()
    writer = csv.writer(buf)
    for row in grid:
        writer.writerow(row)
    return buf.getvalue()


def export_json(annotation: dict) -> str:
    """Serialize annotation to a formatted JSON string."""
    return json.dumps(annotation, indent=2, ensure_ascii=False)


def export_excel(annotation: dict, output_path: str) -> None:
    """Write annotation to an Excel workbook (one sheet per table).

    Requires ``openpyxl``.
    """
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)  # remove default empty sheet

    for table in annotation.get("tables", []):
        ws = wb.create_sheet(title=f"Table {table['table_id']}")
        cells = table.get("cells", [])

        for cell in cells:
            r = cell["row"] + 1      # openpyxl is 1-indexed
            c = cell["col"] + 1
            ws.cell(row=r, column=c, value=cell.get("text", ""))
            if cell["row_span"] > 1 or cell["col_span"] > 1:
                ws.merge_cells(
                    start_row=r, start_column=c,
                    end_row=r + cell["row_span"] - 1,
                    end_column=c + cell["col_span"] - 1,
                )

    wb.save(output_path)
